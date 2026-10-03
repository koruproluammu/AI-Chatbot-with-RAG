"""Offline unit tests: no paid API calls (fake embeddings + fake LLM)."""

from __future__ import annotations

import pytest
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models.fake_chat_models import FakeListChatModel
from pypdf import PdfWriter

from src.document_loader import load_document
from src.rag_chain import (
    NOT_ENOUGH_MESSAGE,
    RAG_PROMPT,
    RAGChain,
    format_context,
)
from src.retriever import Retriever, RetrievedChunk, distance_to_similarity
from src.text_splitter import split_documents
from src.utils import (
    DocumentLoadError,
    EmptyDocumentError,
    RAGChatbotError,
    UnsupportedFileError,
    format_source,
    sanitize_filename,
)
from src.vector_store import VectorStoreManager

TEXT = (
    "Supervised learning trains a model on labeled data. "
    "Examples include spam detection and house price prediction.\n\n"
    "Unsupervised learning finds patterns in unlabeled data, such as clustering."
)


@pytest.fixture
def manager(tmp_path):
    m = VectorStoreManager(tmp_path / "vs", DeterministicFakeEmbedding(size=32))
    chunks = split_documents(
        [Document(page_content=TEXT, metadata={"source": "ml.txt", "page": 2})], 80, 10
    )
    m.build(chunks)
    return m


# ---------------------------------------------------------------- loading
def test_load_txt(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello world")
    docs = load_document(f)
    assert docs[0].metadata == {"source": "a.txt", "doc_type": "txt"}


def test_load_docx(tmp_path):
    f = tmp_path / "a.docx"
    d = DocxDocument()
    d.add_paragraph("Machine learning notes")
    d.save(f)
    assert "Machine learning" in load_document(f)[0].page_content


def test_empty_pdf_raises(tmp_path):
    f = tmp_path / "blank.pdf"
    w = PdfWriter()
    w.add_blank_page(width=72, height=72)
    with open(f, "wb") as fh:
        w.write(fh)
    with pytest.raises(EmptyDocumentError):
        load_document(f)


def test_empty_txt_raises(tmp_path):
    f = tmp_path / "e.txt"
    f.write_text("   \n ")
    with pytest.raises(EmptyDocumentError):
        load_document(f)


def test_corrupted_pdf_raises(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"not a real pdf")
    with pytest.raises(DocumentLoadError):
        load_document(f)


def test_unsupported_extension(tmp_path):
    f = tmp_path / "x.exe"
    f.write_bytes(b"MZ")
    with pytest.raises(UnsupportedFileError):
        load_document(f)


# --------------------------------------------------------------- splitting
def test_split_keeps_metadata_and_size():
    doc = Document(page_content="word " * 500, metadata={"source": "a.pdf", "page": 3})
    chunks = split_documents([doc], 200, 20)
    assert len(chunks) > 1
    assert all(c.metadata["source"] == "a.pdf" and c.metadata["page"] == 3 for c in chunks)
    assert all(len(c.page_content) <= 200 for c in chunks)


def test_split_rejects_bad_overlap():
    with pytest.raises(RAGChatbotError):
        split_documents([Document(page_content="x")], 100, 100)


# --------------------------------------------------------------- retrieval
def test_retrieval_returns_top_k_with_metadata(manager):
    results = Retriever(manager.store, top_k=2).retrieve("What is supervised learning?")
    assert len(results) == 2
    assert results[0].document.metadata["source"] == "ml.txt"


def test_retrieval_rejects_empty_query(manager):
    with pytest.raises(RAGChatbotError):
        Retriever(manager.store).retrieve("   ")


def test_index_persists_and_reloads(manager, tmp_path):
    assert manager.exists()
    fresh = VectorStoreManager(manager.index_dir, DeterministicFakeEmbedding(size=32))
    assert fresh.load() is True
    assert fresh.count() == manager.count()
    assert fresh.indexed_sources() == ["ml.txt"]
    fresh.clear()
    assert not fresh.exists()


def test_distance_to_similarity():
    assert distance_to_similarity(0.0) == 1.0
    assert distance_to_similarity(2.0) == 0.0


# ----------------------------------------------------------------- prompts
def test_prompt_contains_rules_context_and_question():
    text = RAG_PROMPT.format(context="CTX", question="Q?")
    assert "ONLY the provided context" in text
    assert "I couldn't find this information in the uploaded documents." in text
    assert "CTX" in text and "Q?" in text


def test_context_includes_source_labels():
    chunk = RetrievedChunk(
        Document(page_content="body", metadata={"source": "a.pdf", "page": 5}), 0.9
    )
    ctx = format_context([chunk])
    assert "a.pdf — Page 5" in ctx and "body" in ctx


def test_format_source_without_page():
    assert format_source({"source": "n.txt"}) == "n.txt"
    assert sanitize_filename("../../etc/passwd") == "passwd"


# --------------------------------------------------------------- RAG chain
def test_chain_returns_answer_and_sources(manager):
    llm = FakeListChatModel(responses=["Supervised learning uses labeled data."])
    chain = RAGChain(Retriever(manager.store, 2), llm, min_similarity=0.0)
    resp = chain.ask("What is supervised learning?")
    assert resp.answered_from_documents
    assert resp.sources == ["ml.txt — Page 2"]


def test_chain_hides_sources_when_llm_says_not_found(manager):
    llm = FakeListChatModel(
        responses=["I couldn't find this information in the uploaded documents."]
    )
    chain = RAGChain(Retriever(manager.store, 2), llm, min_similarity=0.0)
    resp = chain.ask("Who won the 1998 World Cup?")
    assert resp.sources == [] and not resp.answered_from_documents


def test_chain_skips_llm_when_similarity_too_low(manager):
    llm = FakeListChatModel(responses=["SHOULD NOT BE USED"])
    chain = RAGChain(Retriever(manager.store, 2), llm, min_similarity=1.1)
    resp = chain.ask("anything")
    assert resp.answer == NOT_ENOUGH_MESSAGE


def test_followup_is_rewritten_using_history(manager):
    llm = FakeListChatModel(responses=["What is an example of supervised learning?", "Spam."])
    chain = RAGChain(Retriever(manager.store, 2), llm, min_similarity=0.0)
    history = [
        {"role": "user", "content": "What is supervised learning?"},
        {"role": "assistant", "content": "It uses labeled data."},
    ]
    resp = chain.ask("Give me an example.", history)
    assert resp.standalone_question == "What is an example of supervised learning?"


def test_empty_question_rejected(manager):
    chain = RAGChain(Retriever(manager.store), FakeListChatModel(responses=["x"]))
    with pytest.raises(RAGChatbotError):
        chain.ask("  ")


# ------------------------------------------------- provider / meta handling
def test_gemini_objects_construct_without_network(monkeypatch):
    monkeypatch.setenv("GOOGLE_API_KEY", "AIzaFAKEKEYFORTESTING")
    from src.embeddings import get_embeddings, resolve_embedding_provider
    from src.llm import get_llm

    assert resolve_embedding_provider() == "google"
    assert get_embeddings() is not None
    assert get_llm("google") is not None


def test_missing_google_key_gives_friendly_error(monkeypatch):
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    from src.llm import get_llm
    from src.utils import MissingAPIKeyError

    with pytest.raises(MissingAPIKeyError):
        get_llm("google")


def test_embedding_model_mismatch_is_detected(tmp_path, monkeypatch):
    monkeypatch.setenv("GOOGLE_API_KEY", "AIzaFAKE")
    m = VectorStoreManager(tmp_path)
    m._injected = False
    (tmp_path / "embedding_meta.json").write_text('{"provider": "openai", "model": "x"}')
    (tmp_path / "index.faiss").write_bytes(b"x")
    (tmp_path / "index.pkl").write_bytes(b"x")
    from src.utils import VectorStoreError

    with pytest.raises(VectorStoreError):
        m.load()
