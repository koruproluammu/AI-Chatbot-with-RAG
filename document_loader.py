"""Load PDF / TXT / DOCX files into LangChain Documents with metadata."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

from src.config import MAX_FILE_SIZE_MB, SUPPORTED_EXTENSIONS
from src.utils import (
    DocumentLoadError,
    EmptyDocumentError,
    RAGChatbotError,
    UnsupportedFileError,
)

logger = logging.getLogger(__name__)


def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(str(path))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentLoadError(f"'{path.name}' is password-protected.")
    docs = []
    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        docs.append(
            Document(
                page_content=text,
                metadata={"source": path.name, "page": page_no, "doc_type": "pdf"},
            )
        )
    return docs


def _load_docx(path: Path) -> list[Document]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    # DOCX has no fixed pages, so we deliberately omit the 'page' key.
    return [
        Document(
            page_content="\n".join(parts),
            metadata={"source": path.name, "doc_type": "docx"},
        )
    ]


def _load_txt(path: Path) -> list[Document]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return [Document(page_content=text, metadata={"source": path.name, "doc_type": "txt"})]


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_txt}


def load_document(path: str | Path) -> list[Document]:
    """Load one file. Raises a RAGChatbotError subclass with a friendly message."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileError(
            f"'{path.name}' is not supported. Use: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if not path.is_file():
        raise DocumentLoadError(f"File '{path.name}' was not found.")
    if path.stat().st_size > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise DocumentLoadError(f"'{path.name}' exceeds the {MAX_FILE_SIZE_MB} MB limit.")

    try:
        docs = _LOADERS[ext](path)
    except RAGChatbotError:
        raise
    except Exception as exc:  # corrupted / malformed files
        logger.exception("Failed to parse %s", path.name)
        raise DocumentLoadError(
            f"Could not read '{path.name}'. The file may be corrupted."
        ) from exc

    docs = [d for d in docs if d.page_content.strip()]
    if not docs:
        raise EmptyDocumentError(
            f"'{path.name}' contains no extractable text (scanned PDFs need OCR)."
        )
    return docs
